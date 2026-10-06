"""Build safe, readable XLSX exports for the operator portal."""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from datetime import date, datetime, time
from tempfile import SpooledTemporaryFile
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MOSCOW = ZoneInfo("Europe/Moscow")
_FORMULA_PREFIXES = ("=", "+", "-", "@")


def safe_text(value: Any) -> str:
    """Return valid XML text that Excel cannot interpret as a formula."""
    if value is None:
        return ""
    text = "".join(character for character in str(value) if _is_valid_xml_character(character))
    formula_candidate = text.lstrip(" \t\r\n")
    return "'" + text if formula_candidate.startswith(_FORMULA_PREFIXES) else text


def _is_valid_xml_character(character: str) -> bool:
    codepoint = ord(character)
    return (
        codepoint in (0x09, 0x0A, 0x0D)
        or 0x20 <= codepoint <= 0xD7FF
        or 0xE000 <= codepoint <= 0xFFFD
        or 0x10000 <= codepoint <= 0x10FFFF
    )


def parse_datetime(value: Any, *, utc_to_moscow: bool = False) -> datetime | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value))
    except (TypeError, ValueError):
        return None
    if parsed.tzinfo is not None:
        if utc_to_moscow:
            parsed = parsed.astimezone(MOSCOW)
        parsed = parsed.replace(tzinfo=None)
    return parsed


def parse_date(value: Any) -> date | None:
    try:
        return date.fromisoformat(str(value)) if value else None
    except (TypeError, ValueError):
        return None


def parse_time(value: Any) -> time | None:
    try:
        return time.fromisoformat(str(value)) if value else None
    except (TypeError, ValueError):
        return None


def build_workbook(
    sheet_name: str,
    columns: Sequence[tuple[str, int]],
    rows: Iterable[Sequence[Any]],
    *,
    table_name: str,
) -> SpooledTemporaryFile:
    """Stream a one-sheet workbook into a bounded-memory temporary file."""
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(sheet_name)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 32
    for index, (_, width) in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_cells = []
    for header, _ in columns:
        cell = WriteOnlyCell(sheet, value=header)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_cells.append(cell)
    sheet.append(header_cells)

    body_fill = PatternFill("solid", fgColor="D9EAF7")
    row_count = 0
    for values in rows:
        row_count += 1
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_count % 2 == 0:
                cell.fill = body_fill
            if isinstance(value, datetime):
                cell.number_format = "dd.mm.yyyy hh:mm"
            elif isinstance(value, date):
                cell.number_format = "dd.mm.yyyy"
            elif isinstance(value, time):
                cell.number_format = "hh:mm"
            cells.append(cell)
        sheet.append(cells)

    last_column = get_column_letter(len(columns))
    sheet.auto_filter.ref = f"A1:{last_column}{row_count + 1}"

    # Real Excel Tables need random worksheet access and are not reliable in
    # openpyxl's write-only mode. Auto-filter plus styled/striped rows retain the
    # useful table presentation without keeping the complete export in memory.
    del table_name
    # The response owns this handle and closes it after the download finishes.
    output = SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b")  # noqa: SIM115
    workbook.save(output)
    output.seek(0)
    return output
