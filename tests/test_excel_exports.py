from __future__ import annotations

import os
import uuid
from datetime import date, datetime, time
from io import BytesIO

import pytest
from openpyxl import load_workbook

import database as db
import web


@pytest.fixture()
def export_db(monkeypatch):
    path = os.path.abspath(f".export-test-{uuid.uuid4().hex}.sqlite")
    monkeypatch.setattr(db, "DB_PATH", path)
    monkeypatch.setattr(db, "BOOTSTRAP_ADMIN_PASSWORD", "")
    db.init_db()
    assert db.create_user("admin-export", "a-very-long-password", "Администратор", "admin")[0]
    assert db.create_user("operator-export", "a-very-long-password", "Оператор", "operator")[0]
    yield path
    for suffix in ("", "-wal", "-shm"):
        try:
            os.remove(path + suffix)
        except FileNotFoundError:
            pass


def _login(client, username: str) -> None:
    response = client.post(
        "/login", data={"username": username, "password": "a-very-long-password"},
    )
    assert response.status_code == 302


def _workbook(response):
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["Content-Disposition"].endswith(".xlsx")
    return load_workbook(BytesIO(response.data))


def _assert_layout(sheet, *, rows: bool = True):
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref.startswith("A1:")
    assert sheet.sheet_view.showGridLines is False
    if rows:
        assert len(sheet.tables) == 1


def test_appeals_export_filters_types_formula_safety_and_list_link(export_db):
    conn = db.get_conn()
    conn.executemany(
        """INSERT INTO appeals(ticket_no,ls,channel,category,priority,status,body,
           created_at,closed_at,source) VALUES(?,?,?,?,?,?,?,?,?,?)""",
        [
            ("RSO-1", "+100", "max", "авария", "high", "closed", "=HYPERLINK(\"x\")", "2026-09-02 10:15", "2026-09-02 11:30", "bot"),
            ("RSO-2", "200", "max", "прочее", "normal", "new", "обычный", "2026-09-03 10:15", None, "ai"),
        ],
    )
    conn.commit()
    conn.close()
    client = web.app.test_client()
    assert client.get("/appeals/export.xlsx").status_code == 302
    _login(client, "operator-export")
    page = client.get("/appeals?status=closed&date_from=2026-09-02&date_to=2026-09-02")
    assert "export.xlsx?status=closed" in page.get_data(as_text=True)
    response = client.get("/appeals/export.xlsx?status=closed&date_from=2026-09-02&date_to=2026-09-02")
    workbook = _workbook(response)
    assert workbook.sheetnames == ["Обращения"]
    sheet = workbook.active
    _assert_layout(sheet)
    assert sheet.max_row == 2
    assert sheet["A2"].value == "RSO-1"
    assert sheet["B2"].value == "'+100"
    assert sheet["E2"].value.startswith("'=")
    assert isinstance(sheet["I2"].value, datetime)
    assert sheet["I2"].number_format == "dd.mm.yyyy hh:mm"


def test_appointments_export_empty_and_filtered_rows(export_db):
    conn = db.get_conn()
    branch_id = conn.execute(
        "INSERT INTO branches(name,address) VALUES(?,?)", ("-Филиал", "@Адрес"),
    ).lastrowid
    conn.executemany(
        """INSERT INTO appointments(ls,branch_id,slot_date,slot_time,theme,status,
           channel,chat_id,created_at) VALUES(?,?,?,?,?,?,?,?,?)""",
        [
            ("=10", branch_id, "2026-09-02", "10:30", "+Тема", "active", "max", 1, "2026-09-01 09:00"),
            ("20", branch_id, "2026-09-03", "11:00", "Другая", "visited", "max", 2, "2026-09-01 09:10"),
        ],
    )
    conn.commit()
    conn.close()
    client = web.app.test_client()
    _login(client, "operator-export")
    response = client.get("/appointments/export.xlsx?date_from=2026-09-02&date_to=2026-09-02")
    workbook = _workbook(response)
    assert workbook.sheetnames == ["Записи на приём"]
    sheet = workbook.active
    _assert_layout(sheet)
    assert sheet.max_row == 2
    assert isinstance(sheet["B2"].value, (date, datetime))
    assert sheet["B2"].number_format == "dd.mm.yyyy"
    assert isinstance(sheet["C2"].value, time)
    assert sheet["D2"].value == "'-Филиал"
    assert sheet["E2"].value == "'@Адрес"
    assert sheet["F2"].value == "'=10"
    assert sheet["G2"].value == "'+Тема"
    empty = _workbook(client.get("/appointments/export.xlsx?date=2030-01-01")).active
    _assert_layout(empty, rows=False)
    assert empty.max_row == 1
    assert empty["A1"].value == "Номер записи"


def test_dialog_history_export_is_admin_only_filtered_and_unbounded(export_db):
    conn = db.get_conn()
    operator_id = conn.execute("SELECT id FROM users WHERE username='operator-export'").fetchone()[0]
    rows = []
    for item in range(501):
        rows.append((
            10_000 + item, operator_id, "closed", "@Клиент" if item == 0 else "Клиент",
            "-ЛС", "+Адрес", "=FAQ", "2026-09-02T07:00:00+00:00",
            "2026-09-02T07:01:00+00:00", "2026-09-02T08:00:00+00:00",
            "2026-09-02T09:00:00+00:00",
        ))
    conn.executemany(
        """INSERT INTO operator_dialogs(chat_id,operator_id,status,client_fio,
           client_ls,client_address,faq_context,created_at,first_assigned_at,
           last_activity_at,closed_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)""", rows,
    )
    conn.commit()
    conn.close()
    client = web.app.test_client()
    _login(client, "operator-export")
    denied = client.get("/operator-chat/history/export.xlsx")
    assert denied.status_code == 302
    assert denied.headers["Location"].endswith("/")
    client.get("/logout")
    _login(client, "admin-export")
    page = client.get(
        f"/operator-chat/history?status=closed&operator_id={operator_id}&date_from=2026-09-02&date_to=2026-09-02"
    )
    assert page.status_code == 200
    assert "history/export.xlsx?status=closed" in page.get_data(as_text=True)
    response = client.get(
        f"/operator-chat/history/export.xlsx?status=closed&operator_id={operator_id}&date_from=2026-09-02&date_to=2026-09-02"
    )
    workbook = _workbook(response)
    assert workbook.sheetnames == ["История диалогов"]
    sheet = workbook.active
    _assert_layout(sheet)
    assert sheet.max_row == 502
    assert sheet["B502"].value == "'@Клиент"
    assert sheet["C2"].value == "'-ЛС"
    assert sheet["D2"].value == "'+Адрес"
    assert sheet["K2"].value == "'=FAQ"
    assert sheet["G2"].value.hour == 10  # UTC timestamp is exported in Moscow time.


def test_analytics_has_three_period_scoped_exports(export_db):
    client = web.app.test_client()
    _login(client, "admin-export")
    body = client.get("/analytics?from=2026-09-01&to=2026-09-02").get_data(as_text=True)
    assert "/appeals/export.xlsx?date_from=2026-09-01&amp;date_to=2026-09-02" in body
    assert "/appointments/export.xlsx?date_from=2026-09-01&amp;date_to=2026-09-02" in body
    assert "/operator-chat/history/export.xlsx?date_from=2026-09-01&amp;date_to=2026-09-02" in body
